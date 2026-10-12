import os
import uuid
import re
import time
import json
from typing import List, Dict, Any, Tuple
from pypdf import PdfReader
from pptx import Presentation
from .db import insert_chunks, update_document_status
from .ai_service import get_gemini_client, GEMINI_MODEL

def chunk_text(text: str, max_words: int = 180, overlap: int = 30) -> List[str]:
    """Divides long text into readable chunks with slight overlap."""
    words = text.split()
    if not words:
        return []
    if len(words) <= max_words:
        return [text.strip()]

    chunks = []
    i = 0
    while i < len(words):
        chunk_words = words[i:i + max_words]
        chunks.append(" ".join(chunk_words).strip())
        if i + max_words >= len(words):
            break
        i += (max_words - overlap)
    return chunks

def extract_keywords(text: str, top_k: int = 5) -> List[str]:
    """Extract salient technical keywords from text snippet."""
    stop_words = {
        "the", "and", "is", "in", "to", "of", "a", "with", "for", "as", "on", "that", "this",
        "by", "an", "be", "are", "from", "at", "it", "or", "which", "can", "has", "we", "not"
    }
    words = re.findall(r'\b[a-zA-Z0-9_\-]{3,}\b', text.lower())
    freq = {}
    for w in words:
        if w not in stop_words and not w.isdigit():
            freq[w] = freq.get(w, 0) + 1
    sorted_words = sorted(freq.items(), key=lambda x: x[1], reverse=True)
    return [w[0] for w in sorted_words[:top_k]]

def extract_pdf(file_path: str) -> List[Tuple[str, str, int]]:
    """
    Extracts text from PDF, returning a list of tuples: (content, location, page_number)
    """
    results = []
    try:
        reader = PdfReader(file_path)
    except Exception as e:
        raise ValueError(f"Corrupted or invalid PDF file: {str(e)}")

    if len(reader.pages) == 0:
        raise ValueError("PDF document contains 0 pages.")

    total_extracted_chars = 0
    for page_idx, page in enumerate(reader.pages):
        page_num = page_idx + 1
        page_text = page.extract_text() or ""
        page_text = page_text.strip()
        if page_text:
            total_extracted_chars += len(page_text)
            chunks = chunk_text(page_text, max_words=180, overlap=30)
            for sub_idx, chunk in enumerate(chunks):
                loc = f"Page {page_num}" if len(chunks) == 1 else f"Page {page_num}, §{sub_idx + 1}"
                results.append((chunk, loc, page_num))

    if total_extracted_chars == 0:
        raise ValueError("No extractable text found in PDF. The document may be image-only (scanned) or encrypted.")

    return results

def extract_pptx(file_path: str) -> List[Tuple[str, str, int]]:
    """
    Extracts text from PPTX, returning a list of tuples: (content, location, slide_number)
    """
    results = []
    try:
        prs = Presentation(file_path)
    except Exception as e:
        raise ValueError(f"Corrupted or invalid PowerPoint presentation: {str(e)}")

    if len(prs.slides) == 0:
        raise ValueError("PowerPoint presentation contains 0 slides.")

    total_extracted_chars = 0
    for slide_idx, slide in enumerate(prs.slides):
        slide_num = slide_idx + 1
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    slide_texts.append(text)
        
        slide_content = "\n".join(slide_texts).strip()
        if slide_content:
            total_extracted_chars += len(slide_content)
            chunks = chunk_text(slide_content, max_words=180, overlap=25)
            for sub_idx, chunk in enumerate(chunks):
                loc = f"Slide {slide_num}" if len(chunks) == 1 else f"Slide {slide_num}, Part {sub_idx + 1}"
                results.append((chunk, loc, slide_num))

    if total_extracted_chars == 0:
        raise ValueError("No extractable text found in presentation slides.")

    return results

def extract_txt(file_path: str) -> List[Tuple[str, str, int]]:
    """
    Extracts text from TXT / Markdown file, returning (content, location, line_or_section)
    """
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            full_text = f.read()
    except Exception as e:
        raise ValueError(f"Could not read text file: {str(e)}")

    full_text = full_text.strip()
    if not full_text:
        raise ValueError("Text file is empty.")

    results = []
    lines = full_text.splitlines()
    paragraphs = re.split(r'\n\s*\n', full_text)
    
    current_line = 1
    for p_idx, para in enumerate(paragraphs):
        para_clean = para.strip()
        if not para_clean:
            continue
        para_lines_count = len(para_clean.splitlines())
        end_line = current_line + para_lines_count
        chunks = chunk_text(para_clean, max_words=180, overlap=30)
        for c_idx, chunk in enumerate(chunks):
            loc = f"Lines {current_line}-{end_line}" if len(chunks) == 1 else f"Lines {current_line}-{end_line}, Part {c_idx + 1}"
            results.append((chunk, loc, current_line))
        current_line = end_line + 1

    return results

def parse_timestamp_seconds(ts_str: str) -> int:
    """Parses mm:ss or hh:mm:ss string to integer seconds."""
    try:
        parts = ts_str.strip().split(":")
        if len(parts) == 2:
            return int(parts[0]) * 60 + int(parts[1])
        elif len(parts) == 3:
            return int(parts[0]) * 3600 + int(parts[1]) * 60 + int(parts[2])
        return int(float(ts_str))
    except Exception:
        return 0

def extract_video(file_path: str, doc_id: str) -> Tuple[List[Tuple[str, str, int]], str]:
    """
    Extracts educational concepts and timestamped segments from lecture video (MP4/WebM)
    using the Gemini File API.
    Returns:
        List of (content, location, start_seconds) tuples and overall video summary.
    """
    client = get_gemini_client()
    if not client:
        raise ValueError("GEMINI_API_KEY is required to process and analyze lecture video files. Please configure backend/.env.")

    if not os.path.exists(file_path):
        raise ValueError(f"Video file not found at {file_path}")

    file_size_bytes = os.path.getsize(file_path)
    if file_size_bytes == 0:
        raise ValueError("Video file is empty (0 bytes).")
    if file_size_bytes > 100 * 1024 * 1024:
        raise ValueError("Video file exceeds the 100MB limit for hackathon video processing.")

    # 1. Upload to Gemini Files API
    try:
        remote_file = client.files.upload(file=file_path)
    except Exception as e:
        raise ValueError(f"Failed to upload video to Gemini Files API: {str(e)}")

    # 2. Wait for remote video processing to complete
    max_wait = 180  # 3 minutes
    start_time = time.time()
    try:
        remote_file = client.files.get(name=remote_file.name)
        while remote_file.state.name == "PROCESSING":
            if time.time() - start_time > max_wait:
                raise ValueError("Video processing timed out on remote AI server (exceeded 180s).")
            time.sleep(3)
            remote_file = client.files.get(name=remote_file.name)

        if remote_file.state.name != "ACTIVE":
            raise ValueError(f"Remote video processing ended with status: {remote_file.state.name}")
    except Exception as e:
        try:
            client.files.delete(name=remote_file.name)
        except Exception:
            pass
        raise ValueError(f"Error while waiting for video processing: {str(e)}")

    # 3. Analyze video with Gemini
    prompt = """Analyze this educational lecture video in detail.
Break down the video into chronological, timestamped concept segments where specific definitions, equations, mechanisms, or pedagogical explanations occur.
For every distinct concept segment, provide:
1. start_time: format mm:ss (e.g. "00:00", "01:24")
2. end_time: format mm:ss (e.g. "01:15", "02:45")
3. subtopic: concise title of the concept taught in this timestamp range
4. content: a comprehensive explanation (at least 2-4 sentences) of what the instructor/video explains, presents, or derives in this timestamp range. Include any formulas or definitions visible or spoken.
5. keywords: list of 3-6 key technical terms

Return a JSON object strictly matching this schema:
{
  "summary": "Overall academic summary of the lecture video",
  "topic": "Academic subject of the video",
  "segments": [
    {
      "start_time": "00:00",
      "end_time": "01:15",
      "subtopic": "Concept Name",
      "content": "Explanation...",
      "keywords": ["term1", "term2"]
    }
  ]
}
"""

    response_text = None
    last_error = None

    try:
        models_to_try = [GEMINI_MODEL, "gemini-flash-lite-latest", "gemini-flash-latest"]
        for model_name in models_to_try:
            try:
                resp = client.models.generate_content(
                    model=model_name,
                    contents=[remote_file, prompt],
                    config={
                        "response_mime_type": "application/json",
                        "temperature": 0.2
                    }
                )
                if resp.text and resp.text.strip():
                    response_text = resp.text.strip()
                    break
            except Exception as ex:
                last_error = ex
                continue

        if not response_text:
            raise ValueError(f"Model failed to generate analysis of video: {str(last_error)}")

    finally:
        # Clean up temporary remote file
        try:
            client.files.delete(name=remote_file.name)
        except Exception:
            pass

    # 4. Parse the structured analysis
    try:
        cleaned_json = re.sub(r'^```json\s*', '', response_text)
        cleaned_json = re.sub(r'\s*```$', '', cleaned_json)
        data = json.loads(cleaned_json)
    except json.JSONDecodeError as je:
        raise ValueError(f"Invalid JSON returned during video analysis: {str(je)}")

    summary = data.get("summary", "Lecture video covering educational concepts.")
    raw_segments = data.get("segments", [])
    if not raw_segments:
        raise ValueError("Video analysis returned 0 timestamped concept segments.")

    results = []
    for seg in raw_segments:
        st = seg.get("start_time") or seg.get("start") or "00:00"
        et = seg.get("end_time") or seg.get("end") or "01:00"
        loc = f"Timestamp {st} - {et}"
        subtopic = seg.get("subtopic") or seg.get("label") or seg.get("title") or "Lecture Segment"
        raw_c = seg.get("content") or seg.get("description") or ""
        content_text = f"[{subtopic}] {raw_c}".strip() if raw_c else f"[{subtopic}] Detailed instruction on {subtopic} in lecture video."
        sec = parse_timestamp_seconds(st)
        results.append((content_text, loc, sec))

    return results, summary

def process_uploaded_document(doc_id: str, file_path: str, file_type: str) -> None:
    """
    Synchronous processing of uploaded file.
    Supports PDF, PPTX, TXT, and lecture video (MP4, WebM).
    Updates document status to 'ready' on success or 'error' on failure with specific error message.
    """
    try:
        norm_type = file_type.lower().strip()
        summary_preview = ""

        if norm_type == "pdf":
            extracted = extract_pdf(file_path)
            summary_preview = extracted[0][0][:280] + ("..." if len(extracted[0][0]) > 280 else "")
        elif norm_type in ["pptx", "ppt"]:
            extracted = extract_pptx(file_path)
            summary_preview = extracted[0][0][:280] + ("..." if len(extracted[0][0]) > 280 else "")
        elif norm_type in ["txt", "text", "md"]:
            extracted = extract_txt(file_path)
            summary_preview = extracted[0][0][:280] + ("..." if len(extracted[0][0]) > 280 else "")
        elif norm_type in ["mp4", "webm", "video", "mov", "mkv"]:
            extracted, video_summary = extract_video(file_path, doc_id)
            summary_preview = video_summary[:280] + ("..." if len(video_summary) > 280 else "")
        else:
            raise ValueError(f"Unsupported file format '{file_type}'. Supported: PDF, PPTX, TXT, and Video (MP4, WebM).")

        if not extracted:
            raise ValueError("Extraction produced zero searchable text chunks.")

        chunks_to_insert = []
        for idx, (content, location, page_or_slide) in enumerate(extracted):
            chunk_id = f"chunk-{doc_id}-{idx + 1}"
            keywords = extract_keywords(content)
            chunks_to_insert.append({
                "id": chunk_id,
                "document_id": doc_id,
                "chunk_index": idx + 1,
                "location": location,
                "page_or_slide": page_or_slide,
                "content": content,
                "keywords": keywords
            })

        insert_chunks(chunks_to_insert)

        update_document_status(
            doc_id=doc_id,
            status="ready",
            chunk_count=len(chunks_to_insert),
            summary=summary_preview
        )

    except Exception as e:
        error_msg = str(e)
        update_document_status(
            doc_id=doc_id,
            status="error",
            error_message=error_msg,
            chunk_count=0
        )
        raise e
